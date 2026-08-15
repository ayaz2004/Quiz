import json
import openpyxl
from datetime import date, datetime

WORKBOOK_PATH = r'D:\scholarship database.xlsx'
OUTPUT_PATH = r'.\scripts\scholarshipWorkbook.json'

wb = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
ws = wb['Scholarships']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
rows = []

for row in ws.iter_rows(min_row=2, values_only=True):
    obj = {}
    for i, key in enumerate(headers):
        value = row[i] if i < len(row) else None
        if value is None:
            obj[key] = None
        elif isinstance(value, (datetime, date)):
            obj[key] = value.isoformat()
        elif isinstance(value, str):
            obj[key] = value.strip()
        else:
            obj[key] = value
    rows.append(obj)

with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
    json.dump(rows, f, ensure_ascii=False)

print(f'Exported {len(rows)} scholarship rows to {OUTPUT_PATH}')
