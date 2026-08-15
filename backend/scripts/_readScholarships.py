import json
import openpyxl
from datetime import date, datetime

WORKBOOK_PATH = r'D:\scholarship database.xlsx'

wb = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
ws = wb['Scholarships']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
rows = []

for row in ws.iter_rows(min_row=2, values_only=True):
    obj = {}
    for i, key in enumerate(headers):
        value = row[i] if i < len(row) else None
        if value is None:
            obj[key] = None
        elif isinstance(value, (datetime, date)):
            obj[key] = value.isoformat()
        else:
            obj[key] = str(value).strip() if isinstance(value, str) else value
    rows.append(obj)

print(json.dumps(rows, ensure_ascii=False))
